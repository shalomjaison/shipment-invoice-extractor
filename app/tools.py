from clfl_core_library import DriveManager, extract_year_from_shipment, SheetsManager
from google.genai import types
import os
import json
from concurrent.futures import ThreadPoolExecutor
from app.utils import get_credentials, get_gemini_client

_drive_manager = None
_sheets_manager = None


def get_drive_manager() -> DriveManager:
    global _drive_manager
    if _drive_manager is None:
        _drive_manager = DriveManager(get_credentials())
    return _drive_manager

def get_sheets_manager() -> SheetsManager:
    global _sheets_manager
    if _sheets_manager is None:
        _sheets_manager = SheetsManager(get_credentials())
    return _sheets_manager


# ===== DRIVE TOOLS =====

def find_shipment_folder(shipment_number: str) -> dict:
    """
    Finds the Google Drive folder for a shipment by its number.
    Returns {id, name, drive_id} if found, raises exception if not found.
    """
    drive_manager = get_drive_manager()
    result = drive_manager.find_shipment_folder(shipment_number)
    if not result:
        raise Exception(f"Shipment folder not found: {shipment_number}")
    
    # Add drive_id to the result
    year = extract_year_from_shipment(shipment_number)
    drive_id = drive_manager.get_shared_drive_by_year(year)
    result['drive_id'] = drive_id
    
    return result

def list_folder_files(folder_id: str, drive_id: str) -> list[dict]:
    """
    Lists all files in a folder on a shared drive.
    Returns list of {id, name, mimeType}
    """
    return get_drive_manager().list_shipment_files(folder_id, drive_id)

# def download_file(file_id: str) -> str:
#     """
#     Downloads file content from Google Drive and returns as base64 string.
#     Use this before sending files to Gemini for analysis.
#     """
#     content_bytes = get_drive_manager().download_file_content(file_id)
#     return base64.b64encode(content_bytes).decode('utf-8')

def move_file_to_folder(file_id: str, folder_id: str) -> dict:
    """
    Moves a file into a specified folder.
    Returns {success: True, file_id, new_parent}
    """
    return get_drive_manager().move_file_to_folder(file_id, folder_id)


# ===== SHEETS TOOLS =====

def create_spreadsheet(title: str, folder_id: str = None) -> dict:
    """
    Creates a new Google Spreadsheet via Drive API (bypasses Sheets API permission restriction).
    If folder_id provided, places it directly in that folder.
    Returns {spreadsheet_id, spreadsheet_url}
    """
    
    return get_drive_manager().create_spreadsheet(title, folder_id)

def append_rows(spreadsheet_id: str, range: str, values: list) -> dict:
    """
    Appends rows to a spreadsheet range.
    values = [["row1col1", "row1col2"], ["row2col1", "row2col2"]]
    Returns API result with updates info.
    """
    return get_sheets_manager().append_rows(spreadsheet_id, range, values)

def get_sheet_values(spreadsheet_id: str, range: str) -> list:
    """
    Reads values from a spreadsheet range.
    Returns list of rows (empty list if range is empty).
    """
    return get_sheets_manager().get_values(spreadsheet_id, range)

def batch_get_sheet_values(spreadsheet_id: str, ranges: list) -> list:
    """
    Reads values from multiple ranges in one call.
    Returns list of ValueRange objects.
    """
    return get_sheets_manager().batch_get_values(spreadsheet_id, ranges)

# ======= GEMINI EXTRACTION TOOLS =======
def extract_invoice_data(mime_type: str, filename: str, file_id: str) -> dict:
    """
    Extracts invoice data from a file using Gemini Vision. Accepts a file ID and downloads the file content internally.
    
    Args:
        mime_type: MIME type (e.g., 'application/pdf', 'image/jpeg')
        filename: Original filename for context
        file_id: Google Drive file ID
    
    Returns:
        {
            "invoice_number": str,
            "date": str,
            "total_amount": float,
            "vendor_name": str,
            "currency": str,
            "issued_to": str,
            "description": str
        }
    """
    __genai_client = get_gemini_client()
    prompt = f"""
        You are extracting structured invoice metadata.

        Return ONLY valid JSON.
        Do not include markdown fences.
        Do not include explanations.

        Schema:
        {{
        "invoice_number": string | null,
        "date": string | null,
        "total_amount": number | null,
        "vendor_name": string | null,
        "currency": string | null,
        "issued_to": string | null,
        "description": string | null
        }}

        Rules:
        - Extract from the document only.
        - If a field is missing or uncertain, return null.
        - date must be exactly as seen unless a clean ISO date is obvious.
        - total_amount must be numeric only, no currency symbols, no commas unless needed for decimal parsing.
        - currency should be the currency code or symbol shown in the invoice.
        - description should be a short summary of what the invoice is for.
        - If the invoice is not in English, translate it to English.
        - If output is not valid JSON, fix it before returning.
        - Use the filename only as weak context: {filename}
    """

    file_bytes = get_drive_manager().download_file_content(file_id, mime_type)
    print(f"[DEBUG] Downloaded {len(file_bytes)} bytes for file_id={file_id}, mime_type={mime_type}")
    if not file_bytes:
        raise ValueError(f"Downloaded file is empty: {file_id}")
    response = __genai_client.models.generate_content(
        model="gemini-2.5-flash",
        contents=[prompt, types.Part.from_bytes(data=file_bytes, mime_type=mime_type)],
        config=types.GenerateContentConfig(
        response_mime_type="application/json"
        )
    )
    return json.loads(response.text)


def sniff_file_invoice(file_id: str, mime_type: str) -> str | None:
    """
    Download a file from Google Drive and extract a short text sample for triage.

    Used as a cheap first pass before sending a document to an LLM classifier.
    Extracts text based on MIME type:
      - PDF: text from the first page via pdfplumber
      - DOCX: text from the first paragraph via python-docx
      - XLSX: first 10 rows of the active sheet, cells joined with " | "
      - EML (message/rfc822): decoded plain-text body payload
      - Images (image/*): returns None (no text to extract; caller handles via vision)

    Args:
        file_id: Google Drive file ID to download.
        mime_type: MIME type of the file, used to select the extraction strategy.

    Returns:
        Up to 500 characters of extracted text, or None if the file type
        requires vision-based classification (e.g. images).

    Raises:
        ValueError: If the downloaded file content is empty.
    """
    if mime_type.startswith("image/"):
        return None

    file_bytes = get_drive_manager().download_file_content(file_id, mime_type)
    if not file_bytes:
        raise ValueError(f"Downloaded file is empty: {file_id}")

    import io
    if mime_type == "application/pdf":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            if not pdf.pages:
                return None
            first_page = pdf.pages[0]
            text = first_page.extract_text()
            if text is None:
                return None
            print(f"[DEBUG] Sniffed text: {text}")
    elif mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        text = None
        for paragraph in doc.paragraphs:
            paragraph_text = paragraph.text.strip()
            if paragraph_text:
                text = paragraph_text
                break
        if text is None:
            return None
        print(f"[DEBUG] Sniffed text: {text}")
    elif mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes))
        sheet = workbook.active
        extracted_data = []
        for row in sheet.iter_rows(max_row=10, values_only=True):
            row_text = " | ".join(str(cell) for cell in row if cell is not None)
            if row_text:
                extracted_data.append(row_text)
        workbook.close()
        text = "\n".join(extracted_data)
    elif mime_type == "message/rfc822":
        import email
        msg = email.message_from_bytes(file_bytes)
        payload_bytes = None
        if msg.is_multipart():
            for part in msg.walk():
                content_type = part.get_content_type()
                disposition = (part.get("Content-Disposition") or "").lower()
                if content_type == "text/plain" and "attachment" not in disposition:
                    payload_bytes = part.get_payload(decode=True)
                    if payload_bytes:
                        break
        else:
            payload_bytes = msg.get_payload(decode=True)
        if payload_bytes is None:
            return None
        text = payload_bytes.decode("utf-8", errors="replace")
    else:
        return None

    return text[:500]


def classify_excerpt(file_id: str, mime_type: str, text: str | None, user_prompt: str) -> str:
    """
    Classify a document as relevant or skip using Gemini Flash Lite.

    Two modes depending on whether a text excerpt is available:
      - Text mode (text is not None): sends the user prompt + text excerpt to the
        model as plain text. Faster and cheaper.
      - Vision mode (text is None): downloads the raw file bytes and sends them
        as an inline Part alongside the prompt. Used for image-only documents
        where no text could be extracted by sniff_file_invoice.

    The model is instructed to reply with exactly one word: "relevant" or "skip".

    Args:
        file_id: Google Drive file ID (only downloaded in vision mode).
        mime_type: MIME type of the file (used for inline Part in vision mode).
        text: Short text excerpt from sniff_file_invoice, or None for images.
        user_prompt: The original user request describing what files to look for.

    Returns:
        "relevant" if the document matches the user's request, "skip" otherwise.

    Raises:
        ValueError: If vision mode is triggered but the downloaded file is empty
    """
    __genai_client = get_gemini_client()
    if text is not None:
        prompt = f"""
        You are a document classifier for a freight forwarding company with 15 years of experience handling shipment files.

        You will be given a user prompt and a short excerpt from a document. Your job is to determine whether the document is relevant to the user's request.

        User prompt:
        ```{user_prompt}```

        Document excerpt:
        ```{text}```

        Reply with ONLY one word — either "relevant" or "skip". No explanation. No punctuation. No other text.
        """
        response = __genai_client.models.generate_content(
            model="gemini-2.0-flash-lite",
            contents=[prompt],
            config=types.GenerateContentConfig(response_mime_type="text/plain")
        )
    else:
        file_bytes = get_drive_manager().download_file_content(file_id, mime_type)
        if not file_bytes:
            raise ValueError(f"Downloaded file is empty: {file_id}")

        prompt = f"""
        You are a document classifier for a freight forwarding company with 15 years of experience handling shipment files.
        You will be given an image of a document or a scanned document. Your job is to determine whether the document is relevant to the user's request.

        User prompt:
        ```{user_prompt}```

        Reply with ONLY one word — either "relevant" or "skip". No explanation. No punctuation. No other text.
        """
        response = __genai_client.models.generate_content(
            model="gemini-2.0-flash-lite",
            contents=[prompt, types.Part.from_bytes(data=file_bytes, mime_type=mime_type)],
            config=types.GenerateContentConfig(response_mime_type="text/plain")
        )

    normalized = response.text.strip().lower()
    if normalized in {"relevant", "skip"}:
        return normalized
    return "skip"


def triage_file_invoice(file_id: str, mime_type: str, user_prompt: str) -> str:
    """
    Decide whether a file should be extracted, skipped, or recursed into.

    This is the entry point for the pre-extraction triage stage. It runs a
    layered decision pipeline to avoid sending irrelevant or unsupported files
    to the full Gemini extraction model (extract_invoice_data).

    Pipeline:
      1. Hard MIME skip — immediately returns "skip" for archive, video, and
         audio types that can never contain invoice data.
      2. Folder check — returns "recurse" for Google Drive folders so the
         caller can list and triage their contents recursively.
      3. Text sniff — calls sniff_file_invoice to extract a short text sample
         from the file without invoking an LLM.
      4. LLM classify — passes the text sample (or raw image bytes for images)
         to classify_excerpt, which uses Gemini Flash Lite to decide relevance.

    Args:
        file_id: Google Drive file ID.
        mime_type: MIME type used to route through the pipeline.
        user_prompt: The original user request, forwarded to the LLM classifier.

    Returns:
        "relevant" — file should be passed to extract_invoice_data.
        "skip"     — file is not relevant or not supported; ignore it.
        "recurse"  — file is a folder; list its contents and triage each child.
    """
    hard_skip_mime_types = [
        "application/x-rar-compressed",
        "application/x-7z-compressed",
        "video/mp4",
        "video/quicktime",
        "audio/mpeg",
        "audio/wav",
    ]

    if mime_type in hard_skip_mime_types:
        return "skip"
    if mime_type == "application/vnd.google-apps.folder":
        return "recurse"

    text = sniff_file_invoice(file_id, mime_type)
    if text is None and not mime_type.startswith("image/"):
        return "skip"
    return classify_excerpt(file_id, mime_type, text, user_prompt)

def triage_folder_files(files: list[dict], user_prompt: str) -> dict:
    """
    Triage a batch of Drive files concurrently for invoice relevance.

    Each file is evaluated via `triage_file_invoice`, which returns one of:
    - "relevant": file should be sent for extraction
    - "skip": file is not relevant/processable
    - "recurse": file is a folder and its children should be triaged

    Args:
        files: List of Drive file metadata objects that include at least
            `id` and `mimeType`.
        user_prompt: User request used as classification context.

    Returns:
        A mapping of `file_id -> triage_decision`.
    """
    def triage_one(file: dict):
        return triage_file_invoice(file['id'], file['mimeType'], user_prompt)

    with ThreadPoolExecutor(max_workers=10) as executor:
        futures_by_file_id = {
            file["id"]: executor.submit(triage_one, file)
            for file in files
        }

    results: dict[str, str] = {}
    for file in files:
        file_id = file["id"]
        try:
            results[file_id] = futures_by_file_id[file_id].result()
        except Exception as e:
            print(f"[WARN] Triage failed for file_id={file_id}: {type(e).__name__}: {e}")
            results[file_id] = "error"

    return results

